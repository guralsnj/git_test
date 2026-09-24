import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree

SK = "/root/.claude/skills/synced/5dd94d0b-87b7-4407-be58-13dace0ecad4_8e7c9b18-ec18-4b4c-8b20-f72a48e99518/steel-king-ppt-template/assets/template.pptx"
GREEN = RGBColor(0x3C, 0x7D, 0x4D); STEEL = RGBColor(0x5B, 0x66, 0x70); INK = RGBColor(0x15, 0x18, 0x1A)
HAIR = RGBColor(0xE2, 0xE6, 0xE3); TINT = RGBColor(0xF2, 0xF5, 0xF2); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
L = dict(COVER_MASCOT=15, TITLE_CONTENT=4, TWO_CONTENT=5, THREE_CARDS=6)

# ---------- helpers ----------
def ph(slide, idx):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == idx:
            return sh
    raise KeyError(idx)

def fill(slide, idx, text):
    p = ph(slide, idx).text_frame.paragraphs[0]
    if p.runs: p.runs[0].text = text
    else: p.add_run().text = text

def fill_lines(slide, idx, lines, size=None):
    tf = ph(slide, idx).text_frame
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    p0 = tf.paragraphs[0]
    if p0.runs: p0.runs[0].text = lines[0]
    else: p0.add_run().text = lines[0]
    for r in p0.runs[1:]: r._r.getparent().remove(r._r)
    for line in lines[1:]:
        p = tf.add_paragraph(); p.add_run().text = line
    if size:
        for p in tf.paragraphs:
            for r in p.runs: r.font.size = Pt(size)

def drop(slide, idx):
    sh = ph(slide, idx); sh._element.getparent().remove(sh._element)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def textbox(slide, x, y, w, h, text, size=13, bold=False, color=INK, fill_tint=False, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.10 if fill_tint else 0)
    tf.margin_top = tf.margin_bottom = Inches(0.06 if fill_tint else 0)
    if fill_tint:
        tb.fill.solid(); tb.fill.fore_color.rgb = TINT
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def heading(slide, x, y, w, text):
    return textbox(slide, x, y, w, 0.26, text.upper(), size=11, bold=True, color=GREEN)

def _border(cell, color="E2E6E3", width=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)): tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag), w=str(width), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")

def table(slide, x, y, w, col_w, header, rows, row_h=0.30, size=10, body_size=None, header_h=None):
    body_size = body_size or size
    header_h = header_h or row_h
    n = len(rows) + 1
    gf = slide.shapes.add_table(n, len(header), Inches(x), Inches(y), Inches(w), Inches(header_h + row_h * len(rows)))
    t = gf.table
    tblPr = t._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for st in tblPr.findall(qn("a:tableStyleId")): tblPr.remove(st)
    for i, cw in enumerate(col_w): t.columns[i].width = Inches(cw)
    t.rows[0].height = Inches(header_h)
    for r in range(1, n): t.rows[r].height = Inches(row_h)
    for ci, txt in enumerate(header):
        c = t.cell(0, ci); c.fill.solid(); c.fill.fore_color.rgb = GREEN
        c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = txt
        r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE
        _border(c, "3C7D4D")
    for ri, row in enumerate(rows, start=1):
        for ci, txt in enumerate(row):
            c = t.cell(ri, ci); c.fill.solid(); c.fill.fore_color.rgb = WHITE
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.TOP
            first = ci == 0
            p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = txt
            r.font.name = "Arial"; r.font.size = Pt(body_size); r.font.bold = first
            r.font.color.rgb = INK if first else STEEL
            _border(c)
    return t

def chart(slide, x, y, w, h):
    data = CategoryChartData()
    data.categories = ["Q4 2025", "Q1 2026", "Q2 2026", "Q3 2026"]
    data.add_series("Bookings", (850, 920, 780, 1010))
    data.add_series("Quotes", (2400, 2650, 2100, 2900))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart; ch.has_title = False
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10); ch.legend.font.name = "Arial"
    for s, col in zip(ch.plots[0].series, (GREEN, STEEL)):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = col
    ch.plots[0].gap_width = 55; ch.plots[0].overlap = -10
    ch.value_axis.has_major_gridlines = False; ch.value_axis.visible = False
    ch.category_axis.tick_labels.font.size = Pt(10); ch.category_axis.tick_labels.font.name = "Arial"
    ch.category_axis.format.line.color.rgb = HAIR
    ch.plots[0].has_data_labels = True
    dl = ch.plots[0].data_labels; dl.font.size = Pt(9); dl.font.name = "Arial"; dl.number_format = '$#,##0"K"'; dl.number_format_is_linked = False
    return ch

def strip_watermark(prs):
    """Light layouts carry the shield watermark as a background picture. Replace with plain white."""
    for i in (3, 4, 5, 6, 7, 8, 9, 10, 11, 13):
        layout = prs.slide_layouts[i]
        cSld = layout._element.find(qn("p:cSld"))
        bg = cSld.find(qn("p:bg"))
        if bg is not None:
            cSld.remove(bg)
        new_bg = etree.SubElement(cSld, qn("p:bg")); cSld.insert(0, new_bg)
        bgPr = etree.SubElement(new_bg, qn("p:bgPr"))
        sf = etree.SubElement(bgPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val="FFFFFF")
        etree.SubElement(bgPr, qn("a:effectLst"))

def new_deck():
    prs = Presentation(SK)
    strip_watermark(prs)
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId); del prs.slides._sldIdLst[0]
    return prs

def std_header(slide, n, eyebrow, title, takeaway):
    fill(slide, 100, f"TERRITORY REVIEW  ·  {n} OF 7  ·  {eyebrow}".upper())
    fill(slide, 101, title)
    fill(slide, 102, takeaway)


# ---------- slides ----------
def cover(prs, title, line, who):
    s = prs.slides.add_slide(prs.slide_layouts[L["COVER_MASCOT"]])
    fill(s, 102, "SALES MEETING  ·  TERRITORY REVIEW")
    fill(s, 104, title); fill(s, 105, line); fill(s, 106, who)
    notes(s, "Twelve minutes, hard stop, then questions. Seven slides. Real dealer, account, and job names. "
             "Slides 1 and 2 are pre-filled by Nate. You bring the story.")
    return s

def blank_rows(n, cols, first=None):
    return [[(first[i] if first else "")] + [""] * (cols - 1) for i in range(n)]

def slide_numbers(prs, label, quarters=True):
    s = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    std_header(s, 1, "By the numbers", f"{label} in numbers",
               "[One line on what the numbers say. The story starts here.]")
    drop(s, 103)
    heading(s, 0.80, 2.32, 11.73, "Bookings by quarter, with the company lead time that quarter  ·  pre-filled")
    table(s, 0.80, 2.60, 11.73, [1.60, 2.00, 2.00, 1.60, 2.00, 2.53],
          ["Quarter", "Bookings 2025", "Bookings 2026", "Change", "Quotes 2026 (count)", "Company lead time (wks)"],
          [["Q1", "", "", "", "", ""], ["Q2", "", "", "", "", ""], ["Q3", "", "", "", "", ""],
           ["Q4 to date", "", "", "", "", ""], ["YTD", "", "", "", "", ""]],
          row_h=0.29, size=10, header_h=0.30)
    heading(s, 0.80, 4.82, 11.73, "What the numbers don't show  ·  you fill this in")
    textbox(s, 0.80, 5.10, 11.73, 1.50,
            "[Two or three sentences. What happened in the territory that bookings don't capture. "
            "If bookings didn't come back when lead time did, say why. If they did, say what you did to get them back.]",
            size=12, color=STEEL, fill_tint=True)
    fill(s, 104, "Source: [system], pulled [date]. Quote count, not quote dollars. Lead time is the quoted lead time from the plant that quarter.")
    notes(s, "Slide 1. Pre-filled. Bookings by quarter this year against last year, quote count, and the company lead time that quarter. "
             "Lead times were 30-plus weeks in the first half and have been back to 5 to 10 weeks since summer. "
             "Add a couple of sentences on what the numbers don't show.")
    return s

def slide_dealer_scorecard(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    std_header(s, 2, "Dealer scorecard", "Who carries the territory, and who doesn't",
               "[One line. Where the volume really comes from, and which Excalibur names are on the list for a reason.]")
    drop(s, 103)
    heading(s, 0.80, 2.32, 11.73, "Top five dealers by 2026 bookings  ·  numbers pre-filled, last column is yours")
    table(s, 0.80, 2.60, 11.73, [2.60, 1.00, 1.45, 1.45, 1.00, 4.23],
          ["Dealer", "Excalibur", "2025", "2026 YTD", "GP %", "Who else they sell, and what share we get"],
          [["1.  ", "", "", "", "", ""], ["2.  ", "", "", "", "", ""], ["3.  ", "", "", "", "", ""],
           ["4.  ", "", "", "", "", ""], ["5.  ", "", "", "", "", ""]], row_h=0.29, size=10, header_h=0.30)
    heading(s, 0.80, 4.78, 11.73, "Excaliburs that are flat, down, or at zero  ·  numbers pre-filled, last two columns are yours. Add rows as needed")
    table(s, 0.80, 5.06, 11.73, [2.60, 1.45, 1.45, 1.60, 4.63],
          ["Dealer", "2025", "2026 YTD", "Keep, fix, or drop", "Why, and what fix means"],
          [["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]],
          row_h=0.29, size=10, header_h=0.30)
    fill(s, 104, "Source: [system], pulled [date]. Excalibur list as of [date].")
    notes(s, "Slide 2. Pre-filled from the ranked customer data. Top five dealers with last year next to this year and margin. "
             "Then every Excalibur in the territory that is flat, down, or at zero bookings and quotes this year. "
             "For each top dealer: who else do they sell and what share of their rack business do we get. "
             "For each bottom Excalibur: keep, fix, or drop, and why.")
    return s

def slide_dealers_to_move(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    std_header(s, 3, "Dealers to move", "Who should be in, and who we lost",
               "[One line. The one dealer move that would change the territory the most.]")
    drop(s, 103); drop(s, 104)
    heading(s, 0.80, 2.32, 5.60, "Non-Excaliburs that should be in")
    table(s, 0.80, 2.60, 5.60, [1.70, 1.95, 1.95], ["Dealer", "Why", "What it takes"],
          [["1.  ", "", ""], ["2.  ", "", ""], ["3.  ", "", ""]], row_h=1.10, size=11, header_h=0.40)
    heading(s, 7.13, 2.32, 5.40, "Dealers we lost during the 30-week lead time window")
    table(s, 7.13, 2.60, 5.40, [1.70, 1.85, 1.85], ["Dealer", "Where they went", "Back? What it takes"],
          [["1.  ", "", ""], ["2.  ", "", ""], ["3.  ", "", ""]], row_h=1.10, size=11, header_h=0.40)
    notes(s, "Slide 3. Left: non-Excalibur dealers doing volume or asking to get in, and what it would take. "
             "Right: dealers who went elsewhere while lead times were 30-plus weeks, where they went, and whether they are back. "
             "If nobody left, say so. That tells us something too.")
    return s

def slide_walmart_people(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    std_header(s, 2, "Walmart", "The Walmart relationship",
               "[One line. Where the relationship stands today and what it turns on.]")
    drop(s, 103); drop(s, 104)
    table(s, 0.80, 2.32, 11.73, [2.40, 2.60, 3.93, 2.80],
          ["Decision maker", "Role, what they own", "What they want from us", "Where we stand with them"],
          [["1.  [Name]", "", "", ""], ["2.  [Name]", "", "", ""], ["3.  [Name]", "", "", ""],
           ["4.  [Name]", "", "", ""], ["5.  [Name]", "", "", ""]], row_h=0.776, size=11, header_h=0.40)
    notes(s, "Slide 2 for Chad. Who the decision makers are at Walmart and what they want from us.")
    return s

def slide_walmart_wallet(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    std_header(s, 3, "Share of wallet", "Where Walmart's rack dollars go",
               "[One line. Our share, who has the rest, and why.]")
    drop(s, 103)
    heading(s, 0.80, 2.32, 11.73, "Share of Walmart rack spend by program  ·  your estimate, say so")
    table(s, 0.80, 2.60, 11.73, [2.73, 1.50, 1.50, 1.50, 1.50, 3.00],
          ["Program or format", "Steel King", "Unarco", "Frazier", "Interlake / other", "Why they get it"],
          [["Ambient RDC", "", "", "", "", ""], ["E-comm / fulfillment", "", "", "", "", ""],
           ["Import centers", "", "", "", "", ""], ["[Other]", "", "", "", "", ""]], row_h=0.29, size=10, header_h=0.30)
    heading(s, 0.80, 4.42, 11.73, "Large projects in the pipe  ·  add rows as needed")
    table(s, 0.80, 4.70, 11.73, [3.20, 2.00, 1.50, 1.80, 3.23],
          ["Project", "Type", "$", "Timing", "Where it stands"],
          [["1.  ", "", "", "", ""], ["2.  ", "", "", "", ""], ["3.  ", "", "", "", ""]],
          row_h=0.29, size=10, header_h=0.30)
    fill(s, 104, "Source: share figures are Chad's estimate unless noted.")
    notes(s, "Slide 3 for Chad. Estimated share of Walmart rack spend by program against Unarco, Frazier, and Interlake, "
             "and the large projects in the pipe with timing.")
    return s

def slide_market_opps(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    std_header(s, 2, "New customers", "The five biggest new customer opportunities",
               "[One line. What these five have in common and why they're winnable.]")
    drop(s, 103); drop(s, 104)
    table(s, 0.80, 2.32, 11.73, [3.00, 4.60, 4.13],
          ["Account or segment", "Why it's the opportunity", "What it would take to win it"],
          [["1.  [Account or segment]", "", ""], ["2.  [Account or segment]", "", ""], ["3.  [Account or segment]", "", ""],
           ["4.  [Account or segment]", "", ""], ["5.  [Account or segment]", "", ""]], row_h=0.776, size=11, header_h=0.40)
    notes(s, "Slide 2 for Mac. The five accounts or segments you see as the biggest new customer opportunities and why.")
    return s

def slide_market_inherited(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    std_header(s, 3, "Inherited accounts", "The accounts I took over",
               "[One line. Which of these are alive, which are gone, and which are worth the time.]")
    drop(s, 103)
    table(s, 0.80, 2.32, 11.73, [2.60, 1.45, 1.45, 3.10, 3.13],
          ["Account", "2025", "2026 YTD", "Where it stands", "Next move"],
          [["1.  ", "", "", "", ""], ["2.  ", "", "", "", ""], ["3.  ", "", "", "", ""],
           ["4.  ", "", "", "", ""], ["5.  ", "", "", "", ""]], row_h=0.776, size=11, header_h=0.40)
    fill(s, 104, "Source: bookings pre-filled from [system]. Status and next move are Mac's.")
    notes(s, "Slide 3 for Mac. The accounts that came over with the role. Numbers pre-filled, status and next move are yours.")
    return s

def slide_wins(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["THREE_CARDS"]])
    std_header(s, 4, "Wins", "Three wins",
               "[One line. The pattern across the three, or the one we should be copying.]")
    fill(s, 106, "Price"); fill(s, 111, "Product"); fill(s, 116, "Speed and support")
    for idx, tag in ((107, "price"), (112, "product"), (117, "speed or support")):
        fill_lines(s, idx, ["Job: [dealer, account, job name]",
                            "Size: [$ and product]",
                            f"Why we won on {tag}: [ ]",
                            "Repeatable: [yes or no, and why]"])
    drop(s, 118)
    notes(s, "Slide 4. Three wins. One on price, one on product, one on speed and support. Why we won, and whether it's repeatable.")
    return s

def slide_losses(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["THREE_CARDS"]])
    std_header(s, 5, "Losses", "Three losses",
               "[One line. Which loss hurt most and whether we'd lose it again today.]")
    fill(s, 106, "Price"); fill(s, 111, "Product"); fill(s, 116, "Speed and support")
    fill_lines(s, 107, ["Job: [dealer, account, job]",
                        "Who won: [competitor]",
                        "Their price: [$ or unknown]",
                        "Lead time: ours / theirs [wks]",
                        "Why: [ ]",
                        "Lose it again today: [yes/no]"], size=14)
    fill_lines(s, 112, ["Job: [dealer, account, job]",
                        "Who won: [competitor]",
                        "Why: [what they had, we didn't]",
                        "Lead time: ours / theirs [wks]",
                        "Lose it again today: [yes/no]",
                        "No product loss? Say so here."], size=14)
    fill_lines(s, 117, ["Job: [dealer, account, job]",
                        "Who won: [competitor]",
                        "Lead time: ours / theirs [wks]",
                        "Why: [quoting or service]",
                        "Lose it again today: [yes/no]",
                        "No speed loss? Say so here."], size=14)
    drop(s, 118)
    notes(s, "Slide 5. Three losses. One on price, one on product, one on speed and support. Who won, why, and whether we'd lose it "
             "again today. On the price loss, include the competitor's number if you have it. If you don't, say unknown. "
             "On every loss, the lead time we quoted and the lead time the winner delivered. "
             "If you don't have a loss in one of the three buckets, say so on the slide.")
    return s

def slide_well(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    std_header(s, 6, "Straight talk", "What we do well and what we don't",
               "[One line. The one thing you'd change first.]")
    drop(s, 103); drop(s, 104)
    table(s, 0.80, 2.32, 11.73, [2.33, 4.70, 4.70],
          ["", "What we do well", "What we don't"],
          [["Company", "[One thing]", "[One thing]"],
           ["Marketing", "[One thing]", "[One thing]"],
           ["Sales direction", "[One thing]", "[One thing]"]], row_h=1.293, size=11, body_size=12, header_h=0.40)
    notes(s, "Slide 6. What we do well and what we don't. One of each for the company, for marketing, and for sales direction. "
             "Nate: I mean that last one. I'd rather hear it in the room than not at all.")
    return s

def slide_q4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["THREE_CARDS"]])
    std_header(s, 7, "Q4", "Q4 plan",
               "[One line. What Q4 looks like if these land.]")
    fill(s, 106, "Three projects I'm going after")
    fill_lines(s, 107, ["1. [Job, dealer, $, close date]", "2. [Job, dealer, $, close date]", "3. [Job, dealer, $, close date]"])
    fill(s, 111, "One dealer I'm going to grow")
    fill_lines(s, 112, ["Dealer: [name]", "From: [today's run rate]", "To: [target]", "How: [the specific move]"])
    fill(s, 116, "One thing I'll do differently")
    fill_lines(s, 117, ["What: [ ]", "Why: [what the wins and losses told you]", "What I need from Nate: [ ]"])
    drop(s, 118)
    notes(s, "Slide 7. Q4. Three projects you're going after, one dealer you're going to grow, one thing you'll do differently, "
             "and one thing you need from Nate.")
    return s

def slide_q4_walmart(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["THREE_CARDS"]])
    std_header(s, 7, "2027", "Where Walmart grows in 2027",
               "[One line. The size of the 2027 opportunity and what it turns on.]")
    fill(s, 106, "Where the growth is")
    fill_lines(s, 107, ["1. [DC, format, or program]", "2. [DC, format, or program]", "3. [DC, format, or program]", "Rough size: [$]"])
    fill(s, 111, "What Walmart needs from us to get there")
    fill_lines(s, 112, ["Product: [ ]", "Pricing or terms: [ ]", "Service and lead time: [ ]"])
    fill(s, 116, "What I'll do differently")
    fill_lines(s, 117, ["What: [ ]", "Why: [what the wins and losses told you]", "What I need from Nate: [ ]"])
    drop(s, 118)
    notes(s, "Slide 7 for Chad. Where Walmart grows in 2027.")
    return s

def slide_q4_market(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L["TWO_CONTENT"]])
    std_header(s, 7, "Focus", "Where I'd focus us",
               "[One line. The bet you'd make and what it's worth.]")
    fill(s, 103, "Where I'd focus us")
    fill_lines(s, 105, ["Segment or account: [ ]", "Why now: [ ]", "Rough size of the prize: [$]", "First three moves: [ ]"])
    fill(s, 106, "What I need to get there")
    fill_lines(s, 108, ["From product: [ ]", "From marketing: [ ]", "From pricing: [ ]", "From Nate: [ ]"])
    drop(s, 109)
    notes(s, "Slide 7 for Mac. Where you'd focus us and what you need to get there.")
    return s

# ---------- opener (Nate) ----------
def opener(prs):
    s = prs.slides.add_slide(prs.slide_layouts[7])  # BIG_NUMBERS
    fill(s, 100, "SALES MEETING  ·  BEFORE THE TERRITORY REVIEWS")
    fill(s, 101, "Lead time cost us the first half, not the second")
    fill(s, 102, "The plants weren't full before Walmart. Lead times have been back since summer.")
    fill(s, 104, "30+ to 6"); fill(s, 105, "WEEKS LEAD TIME"); fill(s, 106, "First half of 2026 versus today. New London full through October.")
    fill(s, 108, "19%"); fill(s, 109, "EXCALIBUR SHARE OF 2025"); fill(s, 110, "Historically a third of the business.")
    fill(s, 112, "-50%"); fill(s, 113, "EXCALIBUR SALES, 2025 VS 2024"); fill(s, 114, "Some Excaliburs at zero bookings and quotes in Q3.")
    fill(s, 115, "Today is about what each territory did when lead time came back, and which dealers are carrying us.")
    fill(s, 116, "Source: figures from September leadership meetings. Verify against Don's report and the ranked customer data before presenting.")
    notes(s, "Say this once at the top. Then the reps present. Numbers are from the Aug 31 EC L10, the Sept 2 conversation with Don, "
             "and the Aug 31 1:1 with Brian. Verify before showing.")
    s2 = prs.slides.add_slide(prs.slide_layouts[L["TITLE_CONTENT"]])
    fill(s2, 100, "SALES MEETING  ·  BEFORE THE TERRITORY REVIEWS")
    fill(s2, 101, "Bookings by month against lead time")
    fill(s2, 102, "[One line. Where bookings went when lead time went out, and what they did when it came back.]")
    drop(s2, 103)
    table(s2, 0.80, 2.32, 11.73, [2.20, 2.40, 2.40, 2.00, 2.73],
          ["Month", "Bookings 2025", "Bookings 2026", "Change", "Lead time (wks)"],
          [[m, "", "", "", ""] for m in ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "YTD"]],
          row_h=0.29, size=10, header_h=0.30)
    fill(s2, 104, "Source: [system], pulled [date]. Company-wide, manufactured bookings.")
    notes(s2, "Company-wide bookings by month with the lead time that month. Pre-fill from Don's monthly activity report.")
    return prs

# ---------- build ----------
def build(kind, out):
    prs = new_deck()
    if kind == "rsm":
        cover(prs, "[Territory name]", "The territory the way I see it. Who carries it, where we win, where we lose, what would help.",
              "[Your name]  |  Sales Meeting  |  [Month] 2026")
        slide_numbers(prs, "[Territory]"); slide_dealer_scorecard(prs); slide_dealers_to_move(prs)
        slide_wins(prs); slide_losses(prs); slide_well(prs); slide_q4(prs)
    elif kind == "walmart":
        cover(prs, "Walmart", "The account the way I see it. Who matters, where we win, where we lose, where it grows.",
              "Chad  |  Sales Meeting  |  [Month] 2026")
        slide_numbers(prs, "Walmart"); slide_walmart_people(prs); slide_walmart_wallet(prs)
        slide_wins(prs); slide_losses(prs); slide_well(prs); slide_q4_walmart(prs)
    elif kind == "market":
        cover(prs, "The market", "The market the way I see it. Who matters, where we win, where we lose, where to focus.",
              "Mac  |  Sales Meeting  |  [Month] 2026")
        slide_numbers(prs, "National accounts"); slide_market_opps(prs); slide_market_inherited(prs)
        slide_wins(prs); slide_losses(prs); slide_well(prs); slide_q4_market(prs)
    elif kind == "opener":
        opener(prs)
    prs.save(out)
    print("saved", out)

if __name__ == "__main__":
    outdir = sys.argv[1]
    build("rsm", f"{outdir}/Territory Review Template - RSM.pptx")
    build("walmart", f"{outdir}/Territory Review Template - Walmart (Chad).pptx")
    build("market", f"{outdir}/Territory Review Template - Market (Mac).pptx")
    build("opener", f"{outdir}/Sales Meeting Opener - Nate.pptx")
